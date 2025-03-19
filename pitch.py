import numpy as np
import sys
import pprint
import openpyxl
from intervaltree import IntervalTree

np.set_printoptions(threshold=sys.maxsize)

# Interval tree of frequencies mapping to pitch
def create_tree():
    # Frequency intervals and map of corresponding pitches
    ivs = [
    #Octave 2                     3                    4                    5                    6                # Note
            (63.6,67.4,"C2"),     (127.1,134.7,"C3"),  (254.3,269.4,"C4"),  (508.6,538.8,"C5"),  (1017.1,1077.6,"C6"), # C 
            (67.4,71.4,"C#2"),    (134.7,142.7,"C#3"), (269.4,285.4,"C#4"), (538.8,570.9,"C#5"), (1077.6,1141.7,"C#6"), # C#
            (71.4,75.6,"D2"),     (142.7,151.2,"D3"),  (285.4,302.4,"D4"),  (570.9,604.8,"D5"),  (1141.7,1209.6,"D6"), # D
            (75.6,80.1,"D#2"),    (151.2,160.2,"D#3"), (302.4,320.4,"D#4"), (604.8,640.8,"D#5"), (1209.6,1281.5,"D#6"), # D#
            (80.1,84.9,"E2"),     (160.2,169.7,"E3"),  (320.4,339.4,"E4"),  (640.8,678.9,"E5"),  (1281.5,1357.7,"E6"), # E
            (84.9,89.9,"F2"),     (169.7,179.8,"F3"),  (339.4,359.6,"F4"),  (678.9,719.2,"F5"),  (1357.7,1438.4,"F6"), # F
            (89.9,95.3,"F#2"),    (179.8,190.5,"F#3"), (359.6,381.0,"F#4"), (719.2,762.0,"F#5"), (1438.4,1524.0,"F#6"), # F#
            (95.3,100.9,"G2"),    (190.5,201.8,"G3"),  (381.0,403.7,"G4"),  (762.0,807.3,"G5"),  (1524.0,1614.6,"G6"), # G
            (100.9,106.9,"G#2"),  (201.8,213.8,"G#3"), (403.7,427.7,"G#4"), (807.3,855.3,"G#5"), (1614.6,1710.6,"G#6"), # G#
            (106.9,113.27,"A2"),  (213.8,226.5,"A3"),  (427.7,453.1,"A4"),  (855.3,906.2,"A5"),  (1710.6,1812.3,"A6"), # A
            (113.27,120.0,"A#2"), (226.5,240.0,"A#3"), (453.1,480.0,"A#4"), (906.2,960.1,"A#5"), (1812.3,1920.1,"A#6"), # A#
            (120.0,127.1,"B2"),   (240.0,254.3,"B3"),  (480.0,508.6,"B4"),  (960.1,1017.1,"B5"), (1920.1,2034.3,"B6"),  # B
    
    # Undefined
            (0.0, 63.6, "Undefined"), (2034.3, np.inf, "Undefined")
    ]
    
    tree = IntervalTree.from_tuples(ivs)
    return tree

def proc_freq_data(wb):
    # call the first sheet in the workbook
    ws = wb.active
    # we could also call a worksheet by name for a workbook with multiple sheets
    # ws = wb['sheetname']
    rows = ws.max_row 
    # extract times and pitch frequencies
    ds = 6 # data starts on row 6
    times = []
    freqs = []
    for i in range(ds, rows):
        # zero out any entries with empty frequency values
        time = ws.cell(row=i,column=1).value
        freq = ws.cell(row=i,column=2).value
        
        times.append(time)
        if freq is not None:
            freqs.append(freq)
        else:
            freqs.append(0)

    # get 0.5 sec averages
    times_avg = []
    freqs_avg = []
    step = 38
    for i in range (0, rows-ds, step):
        times_avg.append(times[i])

        freqs_avg.append(np.mean(freqs[i:i+step]))

    # evaluate pitches
    pitches = []
    for freq in freqs_avg:
        res = pitch_tree[freq]
        pitch = list(res).pop()[2]
        pitches.append(pitch)
    
    return times_avg, pitches

# Quick run of pitch analysis

pitch_tree = create_tree()

# hardcoded wb file
wb = openpyxl.load_workbook('data/C Major Scale without Piano Staccoto.xlsx')

times, pitches = proc_freq_data(wb)

pprint.pp("times:\n" + str(times))
pprint.pp(pitches)
